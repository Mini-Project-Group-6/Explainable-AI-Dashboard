"""Text extraction and section segmentation for lesson plans.

Stage 1 of the S1 pipeline (proposal Part D2):
    PDF/DOCX upload -> PyMuPDF / python-docx -> raw text -> section map

The section map feeds ``ingestion/feature_engineer.py`` (see
``docs/FEATURES.md``). Section header keyword lists are a v0.1 draft —
they must be finalised against real GES-format lesson plans from the
tutors (FEATURES.md, open question 2).

Fully offline; deterministic (same file always yields the same output).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

# Canonical section names. "content" covers presentation / core points;
# these five are the main stages used by the time-management features F15/F16.
# Do not add to this tuple to introduce a new section — F15 is a share *of these
# stages*, so extending it silently rescales a trained feature.
MAIN_STAGES = ("introduction", "content", "activities", "assessment", "closure")

# Sections a complete plan is expected to have. Anything here that is absent is
# reported in ``missing_sections`` and drives S2's revision suggestions. Wider
# than MAIN_STAGES on purpose: rpk, resources and differentiation are not timed
# stages but their absence is exactly what the rubric penalises.
EXPECTED_SECTIONS = ("objectives", "rpk", "resources", "differentiation") + MAIN_STAGES

# Ordered: first matching pattern wins, so more specific headers
# (e.g. "learner activities") must come before generic ones.
SECTION_PATTERNS: list[tuple[str, re.Pattern]] = [
    ("objectives", re.compile(
        r"(specific\s+)?(learning\s+)?(objectives?|outcomes?)|by\s+the\s+end\s+of\s+the\s+lesson",
        re.IGNORECASE)),
    ("rpk", re.compile(
        r"\br\.?\s?p\.?\s?k\.?\b|relevant\s+previous\s+knowledge|previous\s+knowledge",
        re.IGNORECASE)),
    # Added in v0.3 for the resources/ICT and attention-to-all-learners
    # criteria, which had no section of their own and so no features.
    ("resources", re.compile(
        r"resources?|materials?|t\.?\s?l\.?\s?m\.?s?\b|teaching\s+aids?|apparatus",
        re.IGNORECASE)),
    ("differentiation", re.compile(
        r"differentiat|inclusi(on|ve)|special\s+educational\s+needs|s\.?e\.?n\.?\b"
        r"|mixed\s+ability|learner\s+diversity|catering",
        re.IGNORECASE)),
    ("introduction", re.compile(
        r"introduction|lesson\s+opening|starter|set\s+induction",
        re.IGNORECASE)),
    ("activities", re.compile(
        r"(teacher.{1,3}learner|learner|pupil|student)?\s*activit(y|ies)|methodolog(y|ies)|methods?\b|procedure",
        re.IGNORECASE)),
    ("content", re.compile(
        r"core\s+points?|content|presentation|subject\s+matter|main\s+ideas?|teaching\s+points?",
        re.IGNORECASE)),
    ("assessment", re.compile(
        r"assessment|evaluation|class\s+exercise|exercise|assignment",
        re.IGNORECASE)),
    ("closure", re.compile(
        r"closure|conclusion|summary|plenary|reflection",
        re.IGNORECASE)),
]

# A header line is short and mostly the keyword itself, e.g.
# "LEARNER ACTIVITIES:" — not a body sentence that merely mentions it.
_MAX_HEADER_WORDS = 8
_MAX_HEADER_CHARS = 60

_DURATION_RE = re.compile(
    r"(?:duration|period|time\s+allocation|time)\s*[:\-]?\s*"
    r"(?:(\d+)\s*(?:hours?|hrs?)\s*)?(\d+)?\s*(?:minutes|mins?)?",
    re.IGNORECASE)


#: Below this many recognised sections, assume the *segmenter* failed rather
#: than the plan. The header keyword lists are still a draft (FEATURES.md open
#: question 2) and a plan laid out in an unanticipated format would otherwise
#: have most of its features silently default to 0 and be scored as a very poor
#: plan. A parsing failure and a bad lesson must not look the same to a tutor.
SECTION_RECOGNITION_FLOOR = 4


@dataclass
class LessonPlanText:
    """Extraction result consumed by the feature engineer."""

    source_path: str
    raw_text: str
    # canonical section name -> section body text ("" if not found)
    sections: dict[str, str] = field(default_factory=dict)
    # expected sections that were not found; surfaced to S2's suggestions panel
    missing_sections: list[str] = field(default_factory=list)
    stated_duration_minutes: Optional[int] = None

    @property
    def sections_found(self) -> int:
        return len(EXPECTED_SECTIONS) - len(self.missing_sections)

    @property
    def format_recognised(self) -> bool:
        """False when the plan looks unparsed rather than poor.

        Consumers must not present scores as findings about teaching quality
        when this is False — see ``format_warning``.
        """
        return self.sections_found >= SECTION_RECOGNITION_FLOOR

    @property
    def format_warning(self) -> Optional[str]:
        """Tutor-facing explanation of a suspected parsing failure, or None."""
        if self.format_recognised:
            return None
        return (
            f"Only {self.sections_found} of {len(EXPECTED_SECTIONS)} expected "
            f"sections were recognised in this document. The scores below are "
            f"probably measuring a layout this tool does not yet read, not the "
            f"quality of the lesson. Check that the plan uses headed sections "
            f"(objectives, introduction/RPK, resources, activities, assessment, "
            f"closure) before acting on the feedback.")


def extract_raw_text(path: str | Path) -> str:
    """Extract plain text from a PDF or DOCX lesson plan."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    raise ValueError(f"Unsupported file type '{suffix}': {path} (expected .pdf or .docx)")


def _extract_pdf(path: Path) -> str:
    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        pages = [page.get_text("text") for page in doc]
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    # Lesson plans are frequently laid out as tables (stage | time | activity):
    # flatten table cells row by row so their text is not lost.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


def _match_header(line: str) -> Optional[str]:
    """Return the canonical section name if *line* looks like a section header.

    A body sentence that merely mentions a keyword ("Exercise to be given
    later.") must NOT count, so a header additionally has to be one of:
    colon-terminated / colon-prefixed ("Evaluation: ..."), ALL CAPS, or
    at most 4 words — and the keyword must match at the start of the line
    (allowing a leading "3." style enumerator).
    """
    stripped = line.strip()
    if not stripped:
        return None
    prefix, colon, _ = stripped.partition(":")
    candidate = (prefix if colon else stripped).strip()
    candidate = re.sub(r"^[\d\.\)\(\s•\-]+", "", candidate)  # "3. EVALUATION" -> "EVALUATION"
    if not candidate:
        return None
    if len(candidate) > _MAX_HEADER_CHARS or len(candidate.split()) > _MAX_HEADER_WORDS:
        return None
    header_like = bool(colon) or candidate.isupper() or len(candidate.split()) <= 4
    if not header_like:
        return None
    for name, pattern in SECTION_PATTERNS:
        if pattern.match(candidate):
            return name
    return None


def segment_sections(raw_text: str) -> tuple[dict[str, str], list[str]]:
    """Split raw text into canonical sections by header keyword lines.

    Lines before the first recognised header form the "header" block
    (school, subject, class, duration, ...). Returns (sections, missing).
    """
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in raw_text.splitlines():
        matched = _match_header(line)
        if matched:
            current = matched
            sections.setdefault(current, [])
            # Keep any text on the header line after a colon ("RPK: learners can add")
            _, _, remainder = line.partition(":")
            if remainder.strip():
                sections[current].append(remainder.strip())
            continue
        sections.setdefault(current, []).append(line)

    joined = {name: "\n".join(lines).strip() for name, lines in sections.items()}
    missing = [name for name in EXPECTED_SECTIONS if not joined.get(name)]
    for name in missing:
        logger.info("Section not found: %s", name)
    return joined, missing


def section_spans(raw_text: str) -> list[tuple[str, int, int]]:
    """(section_name, start_char, end_char) for each section, in document order.

    Same header detection as ``segment_sections``, but keeping character
    offsets so a span of prose can be traced back to the section it came from
    (used by ``explainability/shap_deep.py`` to say *where* in the plan a text
    attribution lands). Offsets index into *raw_text* unchanged.
    """
    spans: list[tuple[str, int, int]] = []
    current, section_start, cursor = "header", 0, 0
    for line in raw_text.splitlines(keepends=True):
        matched = _match_header(line)
        if matched:
            spans.append((current, section_start, cursor))
            current, section_start = matched, cursor
        cursor += len(line)
    spans.append((current, section_start, cursor))
    return [(name, start, end) for name, start, end in spans if end > start]


def section_at(spans: list[tuple[str, int, int]], offset: int) -> str:
    """Canonical section name containing *offset*, or "" if outside them all."""
    for name, start, end in spans:
        if start <= offset < end:
            return name
    return ""


def parse_stated_duration(raw_text: str, header_text: str = "") -> Optional[int]:
    """Stated lesson duration in minutes, from the plan header (e.g. 'Duration: 60 minutes')."""
    search_space = header_text or "\n".join(raw_text.splitlines()[:20])
    for match in _DURATION_RE.finditer(search_space):
        hours, minutes = match.group(1), match.group(2)
        if hours is None and minutes is None:
            continue
        total = int(hours or 0) * 60 + int(minutes or 0)
        if total > 0:
            return total
    return None


def extract(path: str | Path) -> LessonPlanText:
    """Full ingestion for one lesson plan file: text + sections + metadata."""
    raw_text = extract_raw_text(path)
    sections, missing = segment_sections(raw_text)
    duration = parse_stated_duration(raw_text, sections.get("header", ""))
    return LessonPlanText(
        source_path=str(path),
        raw_text=raw_text,
        sections=sections,
        missing_sections=missing,
        stated_duration_minutes=duration,
    )


if __name__ == "__main__":
    import argparse
    import json

    parser = argparse.ArgumentParser(description="Extract text + sections from a lesson plan")
    parser.add_argument("path", help="Path to a .pdf or .docx lesson plan")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO)
    result = extract(args.path)
    print(json.dumps({
        "source_path": result.source_path,
        "stated_duration_minutes": result.stated_duration_minutes,
        "missing_sections": result.missing_sections,
        "sections": {k: v[:200] for k, v in result.sections.items()},
    }, indent=2))
